from openpyxl import load_workbook

from ytps.chunking import Chunk
from ytps.models import Song
from ytps.writers.csv_out import write_csv, write_m3u
from ytps.writers.html_console import write_console
from ytps.writers.xlsx import write_xlsx


def songs(n=30):
    return [Song(position=i, video_id=f"{i:011d}", title=f"Artist {i} - Track {i}",
                 duration="3:30", views=i * 100, category="Party") for i in range(1, n + 1)]


def test_xlsx_has_both_sheets_and_all_rows(tmp_path):
    p = write_xlsx(songs(), tmp_path / "x.xlsx", name="Gym", description="hard")
    wb = load_workbook(p)
    assert wb.sheetnames == ["Playlist", "Build on YouTube"]
    assert wb["Playlist"].max_row == 31
    assert wb["Playlist"].freeze_panes == "C2"


def test_xlsx_rating_and_notes_left_blank_for_the_user(tmp_path):
    wb = load_workbook(write_xlsx(songs(5), tmp_path / "x.xlsx"))
    ws = wb["Playlist"]
    hdr = [c.value for c in ws[1]]
    for col in ("Rating", "Notes"):
        i = hdr.index(col) + 1
        assert all(ws.cell(row=r, column=i).value is None for r in range(2, 7))


def test_csv_roundtrips_through_the_loader(tmp_path):
    from ytps.cli import _load_songs

    p = write_csv(songs(10), tmp_path / "a.csv")
    back = _load_songs(p)
    assert len(back) == 10
    assert back[0].video_id == "00000000001"
    assert back[0].duration_sec == 210


def test_m3u_is_extended_format(tmp_path):
    text = write_m3u(songs(3), tmp_path / "a.m3u", "Gym").read_text()
    lines = text.splitlines()
    assert lines[0] == "#EXTM3U" and lines[1] == "#PLAYLIST:Gym"
    assert lines[2].startswith("#EXTINF:210,")
    assert lines[3].startswith("https://www.youtube.com/watch?v=")


def test_console_is_self_contained_and_theme_aware(tmp_path):
    chunks = [Chunk("Gym", "hard", songs(60)), Chunk("Chill", "easy", songs(10))]
    html = write_console(chunks, tmp_path / "c.html", "Build").read_text()
    assert "<script src=" not in html and "cdn" not in html.lower()   # no external deps
    assert "prefers-color-scheme" in html and 'data-theme="dark"' in html
    assert html.count('class="card"') == 2
    assert html.count('class="chip"') == 3                            # 2 blocks + 1
