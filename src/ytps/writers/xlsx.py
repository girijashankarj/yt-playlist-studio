"""Excel export: a Playlist sheet plus a Build on YouTube sheet with import links."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..models import Song
from ..publish.bulk_links import STEPS, build_links

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
BAND = PatternFill("solid", fgColor="F2F5FA")
LINK = Font(color="0563C1", underline="single")
TITLE = Font(bold=True, size=14, color="1F3864")
SUB = Font(bold=True, size=11, color="1F3864")
EDGE = Border(bottom=Side(style="thin", color="D6DCE4"))

COLS = [
    ("Position", 9), ("Song Name", 34), ("Full Video Title", 52), ("Artist", 28),
    ("Album / Movie", 24), ("Language", 13), ("Genre", 17), ("Sub-Genre", 15),
    ("Category", 12), ("Duration", 10), ("Duration (sec)", 13), ("Views", 14),
    ("Channel", 24), ("Video ID", 14), ("Song Link", 40), ("Rating", 9), ("Notes", 28),
]
CENTER = {1, 6, 7, 8, 9, 10, 11, 16}


def write_xlsx(
    songs: list[Song], path: str | Path, name: str = "", description: str = "",
    source_url: str = "",
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Playlist"

    for i, (head, width) in enumerate(COLS, 1):
        c = ws.cell(row=1, column=i, value=head)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30

    for pos, s in enumerate(songs, 1):
        ws.append([
            pos, s.song_name, s.title, s.artist, s.album, s.language, s.genre,
            s.sub_genre, s.category, s.duration, s.duration_sec, s.views,
            s.channel, s.video_id, s.url, None, None,
        ])
        r = ws.max_row
        for col in range(1, len(COLS) + 1):
            cell = ws.cell(row=r, column=col)
            cell.border = EDGE
            if r % 2 == 0:
                cell.fill = BAND
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=col in (2, 3, 4, 5, 13, 17),
                horizontal="center" if col in CENTER else "general",
            )
        ws.cell(row=r, column=12).number_format = "#,##0"
        lc = ws.cell(row=r, column=15)
        lc.hyperlink, lc.font = s.url, LINK

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{ws.max_row}"

    bs = wb.create_sheet("Build on YouTube")
    for col, w in zip("ABCD", (14, 82, 12, 18), strict=False):
        bs.column_dimensions[col].width = w
    bs.cell(row=1, column=1, value=name or "Playlist").font = TITLE
    total = sum(s.duration_sec for s in songs)
    rows = [
        ("What it is", description), ("Songs", len(songs)),
        ("Runtime", f"{total // 3600}h {total % 3600 // 60:02d}m"),
    ]
    if source_url:
        rows.append(("Source", source_url))
    r = 3
    for k, v in rows:
        bs.cell(row=r, column=1, value=k).font = SUB
        bs.cell(row=r, column=2, value=v)
        r += 1
    if source_url:
        bs.cell(row=r - 1, column=2).hyperlink = source_url
        bs.cell(row=r - 1, column=2).font = LINK

    r += 1
    bs.cell(row=r, column=1, value="How to build it").font = SUB
    for step in STEPS:
        r += 1
        bs.cell(row=r, column=2, value=step)

    r += 2
    for j, h in enumerate(("Block", "Bulk-import link", "Videos", "Positions"), 1):
        c = bs.cell(row=r, column=j, value=h)
        c.fill, c.font = HDR_FILL, HDR_FONT
    for link in build_links(songs):
        r += 1
        bs.cell(row=r, column=1, value=f"Link {link['index']}")
        c = bs.cell(row=r, column=2, value=link["url"])
        c.hyperlink, c.font = link["url"], LINK
        bs.cell(row=r, column=3, value=link["count"])
        bs.cell(row=r, column=4, value=link["positions"])

    r += 2
    bs.cell(row=r, column=1, value="All video IDs in order").font = SUB
    r += 1
    for j, h in enumerate(("Position", "Video ID", "Song Name"), 1):
        c = bs.cell(row=r, column=j, value=h)
        c.fill, c.font = HDR_FILL, HDR_FONT
    for pos, s in enumerate(songs, 1):
        bs.cell(row=r + pos, column=1, value=pos)
        bs.cell(row=r + pos, column=2, value=s.video_id)
        bs.cell(row=r + pos, column=3, value=s.song_name)

    wb.save(path)
    return path


def write_index_xlsx(chunks, path: str | Path, title: str = "Planned playlists") -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Index"
    for col, w in zip("ABCD", (24, 40, 10, 12), strict=False):
        ws.column_dimensions[col].width = w
    ws.cell(row=1, column=1, value=title).font = TITLE
    for j, h in enumerate(("Playlist", "What it is", "Songs", "Runtime"), 1):
        c = ws.cell(row=3, column=j, value=h)
        c.fill, c.font = HDR_FILL, HDR_FONT
    for i, ch in enumerate(chunks):
        r = 4 + i
        ws.cell(row=r, column=1, value=ch.name)
        ws.cell(row=r, column=2, value=ch.description)
        ws.cell(row=r, column=3, value=len(ch.songs)).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=4, value=ch.runtime).alignment = Alignment(horizontal="center")
        for col in range(1, 5):
            ws.cell(row=r, column=col).border = EDGE
    ws.freeze_panes = "A4"
    wb.save(path)
    return path
